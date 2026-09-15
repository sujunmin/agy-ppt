#!/usr/bin/env python3
"""Phase 16.5 final traceability and PPTX Notes tests."""

from __future__ import annotations

import hashlib
import json
import sys
import tempfile
import unittest
from pathlib import Path

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

from PIL import Image  # noqa: E402
from pptx import Presentation  # noqa: E402

from assemble_ppt import create_presentation  # noqa: E402
from phase16_claims import Claim, ContentOrigin, binding_from_unit  # noqa: E402
from phase16_editorial import EditorialReport, EditorialWarning, EDITORIAL_REVIEW_RECOMMENDED  # noqa: E402
from phase16_slide_evidence import MaterialClaim, SlideEvidencePlan  # noqa: E402
from phase16_traceability import (  # noqa: E402
    ERROR_FINAL_TRACEABILITY_INCOMPLETE,
    ERROR_FINAL_TRACEABILITY_INVALID,
    ERROR_STALE_EVIDENCE,
    FinalSlideRecord,
    FinalTraceabilityError,
    build_traceability_report,
    compose_traceable_notes,
    dependency_impact,
    notes_map,
    save_internal_report,
)
from source_grounding import SourceInventory  # noqa: E402


class FinalTraceabilityTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.inv = SourceInventory.initialize(self.root, "final")
        self.digest_a = hashlib.sha256(b"source-a").hexdigest()
        self.digest_b = hashlib.sha256(b"source-b").hexdigest()
        self.inv.add_source("src_a", "pdf", label="Annual report", source_digest=self.digest_a)
        self.inv.add_source("src_b", "image", label="Research chart", source_digest=self.digest_b)
        self.unit_a = self.inv.add_unit("src_a", "page", {"kind": "page", "start": 2, "end": 2}, "HIGH")
        self.unit_b = self.inv.add_unit("src_b", "image", {"kind": "generic", "label": "image:1-of-1"}, "HIGH")

    def tearDown(self):
        self.temp.cleanup()

    def record(self, slide_id="slide_01", extra=(), unsupported=()):
        grounded = Claim.create(self.inv, "Revenue rose 18% in 2025", ContentOrigin.SOURCE_GROUNDED, [binding_from_unit(self.inv, self.unit_a["unit_id"])])
        materials = [MaterialClaim(grounded), *[MaterialClaim(claim) for claim in extra]]
        plan = SlideEvidencePlan.create(slide_id, materials)
        return FinalSlideRecord(plan, "Revenue rose 18% in 2025", "Open with the result.", tuple(unsupported))

    def test_complete_report_counts_provenance_and_evidence(self):
        user = Claim.create(self.inv, "Our target is 20%", ContentOrigin.USER_PROVIDED)
        synthesis = Claim.create(self.inv, "Retention should lead", ContentOrigin.AGY_SYNTHESIS)
        record = self.record(extra=(user, synthesis))
        record = FinalSlideRecord(record.plan, "Revenue rose 18% in 2025 and our target is 20%", record.speaker_notes)
        report = build_traceability_report(self.inv, [record])
        self.assertEqual((report.grounded_claims, report.user_provided_claims, report.agy_synthesis_claims), (1, 1, 1))
        report.require_complete()

    def test_unsupported_claim_fails_final_support_qa(self):
        report = build_traceability_report(self.inv, [self.record(unsupported=("Unsupported source claim",))])
        with self.assertRaises(FinalTraceabilityError) as caught:
            report.require_complete()
        self.assertEqual(caught.exception.error_code, ERROR_FINAL_TRACEABILITY_INCOMPLETE)

    def test_source_change_invalidates_only_dependencies(self):
        unaffected = Claim.create(self.inv, "Image is illustrative", ContentOrigin.USER_PROVIDED)
        other_plan = SlideEvidencePlan.create("slide_02", [MaterialClaim(unaffected)])
        records = [self.record(), FinalSlideRecord(other_plan, "Image is illustrative")]
        impact = dependency_impact(records, {"src_a": hashlib.sha256(b"changed").hexdigest(), "src_b": self.digest_b})
        self.assertEqual(impact.slide_ids, ("slide_01",))

    def test_stale_report_blocks_completion(self):
        report = build_traceability_report(self.inv, [self.record()], current_source_digests={"src_a": "0" * 64})
        with self.assertRaises(FinalTraceabilityError) as caught:
            report.require_complete()
        self.assertEqual(caught.exception.error_code, ERROR_STALE_EVIDENCE)

    def test_inventory_mutation_makes_reference_invalid(self):
        record = self.record()
        self.inv.data["units"].clear()
        with self.assertRaises(FinalTraceabilityError) as caught:
            build_traceability_report(self.inv, [record])
        self.assertEqual(caught.exception.error_code, ERROR_FINAL_TRACEABILITY_INVALID)

    def test_factual_drift_is_checked_at_final_qa(self):
        record = self.record()
        with self.assertRaises(Exception) as caught:
            build_traceability_report(self.inv, [FinalSlideRecord(record.plan, "Revenue rose 28% in 2025")])
        self.assertIn("FACTUAL_DRIFT", caught.exception.error_code)

    def test_editorial_warnings_are_reported_internally(self):
        editorial = EditorialReport(EDITORIAL_REVIEW_RECOMMENDED, (EditorialWarning("COPY_REPETITION", ("slide_01",), "signal"),))
        report = build_traceability_report(self.inv, [self.record()], editorial_report=editorial)
        self.assertEqual(report.editorial_warnings, ("COPY_REPETITION",))

    def test_notes_are_human_readable_and_keep_speaker_copy_first(self):
        notes = compose_traceable_notes(self.record(), self.inv)
        self.assertTrue(notes.startswith("Open with the result."))
        self.assertIn("Annual report — page 2", notes)
        self.assertNotIn("src_a", notes)
        self.assertNotIn("su:", notes)

    def test_user_and_synthesis_provenance_in_notes(self):
        user = Claim.create(self.inv, "User fact", ContentOrigin.USER_PROVIDED)
        synthesis = Claim.create(self.inv, "Recommendation", ContentOrigin.AGY_SYNTHESIS)
        record = self.record(extra=(user, synthesis))
        notes = compose_traceable_notes(record, self.inv)
        self.assertIn("User-provided: User fact", notes)
        self.assertIn("AGY synthesis: Recommendation", notes)

    def test_notes_map_uses_slide_number(self):
        self.assertEqual(list(notes_map([self.record("slide_02")], self.inv)), [2])

    def test_internal_report_persistence_is_deterministic(self):
        report = build_traceability_report(self.inv, [self.record()])
        path = save_internal_report(self.root, report)
        first = path.read_bytes()
        second = save_internal_report(self.root, report).read_bytes()
        self.assertEqual(first, second)
        self.assertEqual(json.loads(first), report.to_dict())

    def test_final_pptx_notes_traceability_and_clean_visible_slide(self):
        image = self.root / "slide_01.png"
        Image.new("RGB", (1600, 900), "white").save(image)
        output = self.root / "final.pptx"
        record = self.record()
        self.assertTrue(create_presentation([str(image)], str(output), speaker_notes=notes_map([record], self.inv)))
        deck = Presentation(output)
        notes = deck.slides[0].notes_slide.notes_text_frame.text
        self.assertIn("Annual report — page 2", notes)
        self.assertEqual(len(deck.slides[0].shapes), 1)
        self.assertNotIn("Traceability", " ".join(shape.text for shape in deck.slides[0].shapes if hasattr(shape, "text")))

    def test_multi_source_binding_remains_complete(self):
        refs = [binding_from_unit(self.inv, self.unit_a["unit_id"]), binding_from_unit(self.inv, self.unit_b["unit_id"])]
        claim = Claim.create(self.inv, "Combined evidence 18%", ContentOrigin.SOURCE_GROUNDED, refs)
        plan = SlideEvidencePlan.create("slide_01", [MaterialClaim(claim)])
        report = build_traceability_report(self.inv, [FinalSlideRecord(plan, "Combined evidence 18%")])
        self.assertEqual(report.evidence_bindings, 2)

    def test_visible_text_is_not_modified_with_citations(self):
        record = self.record()
        build_traceability_report(self.inv, [record])
        self.assertEqual(record.rendered_text, "Revenue rose 18% in 2025")


if __name__ == "__main__":
    unittest.main()
