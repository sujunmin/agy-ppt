#!/usr/bin/env python3
"""Deterministic contracts for v0.6.0 layout and native typography hardening."""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.oxml.ns import qn

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

from phase18_contract import (  # noqa: E402
    DeliveryEditabilityContract,
    EditabilityClass,
    EditabilityEnvelope,
    FontPortability,
    ProductionStrategy,
)
from phase18_hybrid_pptx import (  # noqa: E402
    ChartSeries,
    ChartSpec,
    ChartStyle,
    HybridElement,
    HybridPptxError,
    HybridSlide,
    ImageStyle,
    ShapeStyle,
    TextStyle,
    create_hybrid_presentation,
    text_style_for_role,
)
from phase18_production_plan import (  # noqa: E402
    ElementProductionPlan,
    ElementRole,
    PortabilityRisk,
)
from presentation_layout_grammar import (  # noqa: E402
    DEFAULT_LAYOUT_GRAMMAR,
    LayoutIssue,
    LayoutSeverity,
    audit_layout,
    audit_text_envelope,
    estimated_line_count,
    max_characters_for_box,
)


def plan(element_id: str, role: ElementRole, strategy: ProductionStrategy, content: str):
    return ElementProductionPlan(
        f"ep:{element_id}", element_id, f"semantic:{element_id}", f"narrative:{element_id}",
        content, (), role, PortabilityRisk.LOW,
        DeliveryEditabilityContract(
            EditabilityClass.REPLACEABLE if strategy is ProductionStrategy.NATIVE_IMAGE else EditabilityClass.EDITABLE_PREFERRED,
            strategy, FontPortability.FONT_SAFE,
        ),
    )


class PresentationQualityGrammarTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.photo = self.root / "wide.png"
        Image.new("RGB", (800, 400), (25, 60, 105)).save(self.photo)

    def tearDown(self):
        self.temp.cleanup()

    def render(self, elements):
        path = self.root / "deck.pptx"
        create_hybrid_presentation((HybridSlide("slide_01", tuple(elements)),), str(path))
        return Presentation(path)

    def test_10x5625_coordinate_system_matches_assembly(self):
        deck = self.render((
            HybridElement(plan("title", ElementRole.TITLE, ProductionStrategy.NATIVE_TEXT, "A point"), DEFAULT_LAYOUT_GRAMMAR.title_box()),
        ))
        self.assertAlmostEqual(deck.slide_width / 914400, 10.0, places=3)
        self.assertAlmostEqual(deck.slide_height / 914400, 5.625, places=3)

    def test_three_columns_fit_with_consistent_gaps(self):
        columns = DEFAULT_LAYOUT_GRAMMAR.columns(3)
        self.assertEqual(len(columns), 3)
        self.assertTrue(all(DEFAULT_LAYOUT_GRAMMAR.inside_slide(box) for box in columns))
        gaps = [columns[index + 1].left - (columns[index].left + columns[index].width) for index in range(2)]
        self.assertTrue(all(abs(gap - DEFAULT_LAYOUT_GRAMMAR.column_gap) < 0.001 for gap in gaps))

    def test_out_of_bounds_native_object_fails_before_pptx(self):
        element = HybridElement(
            plan("bad", ElementRole.BODY, ProductionStrategy.NATIVE_TEXT, "Off slide"),
            DEFAULT_LAYOUT_GRAMMAR.content_box(),
        )
        object.__setattr__(element, "box", type(element.box)(9.5, 1, 2, 1))
        with self.assertRaises(HybridPptxError):
            HybridSlide("slide_01", (element,))

    def test_role_typography_has_clear_title_body_kpi_hierarchy(self):
        title = text_style_for_role(ElementRole.TITLE)
        body = text_style_for_role(ElementRole.BODY)
        kpi = text_style_for_role(ElementRole.KPI)
        self.assertGreaterEqual(title.font_size / body.font_size, 1.7)
        self.assertGreater(kpi.font_size, title.font_size)
        self.assertTrue(title.bold)
        self.assertTrue(kpi.bold)

    def test_native_text_has_bounded_margins_and_spacing(self):
        style = text_style_for_role(ElementRole.TITLE)
        deck = self.render((HybridElement(
            plan("title", ElementRole.TITLE, ProductionStrategy.NATIVE_TEXT, "成果不是口號，而是續約率"),
            DEFAULT_LAYOUT_GRAMMAR.title_box(), text_style=style,
        ),))
        frame = deck.slides[0].shapes[0].text_frame
        self.assertEqual(frame.margin_left, 0)
        self.assertEqual(frame.margin_right, 0)
        self.assertAlmostEqual(frame.paragraphs[0].line_spacing, 1.02, places=2)

    def test_cjk_font_is_encoded_separately_from_latin_font(self):
        deck = self.render((HybridElement(
            plan("title", ElementRole.TITLE, ProductionStrategy.NATIVE_TEXT, "續約率帶動成長"),
            DEFAULT_LAYOUT_GRAMMAR.title_box(),
        ),))
        run = deck.slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
        east_asian = run._r.get_or_add_rPr().find(qn("a:ea"))
        self.assertIsNotNone(east_asian)
        self.assertEqual(east_asian.get("typeface"), "Microsoft JhengHei")

    def test_cover_crop_preserves_aspect_without_stretching(self):
        element = HybridElement(
            plan("photo", ElementRole.PHOTO, ProductionStrategy.NATIVE_IMAGE, "photo"),
            type(DEFAULT_LAYOUT_GRAMMAR.content_box())(0.6, 1.4, 3.0, 3.0),
            image_path=str(self.photo), image_style=ImageStyle(crop_mode="cover"),
        )
        deck = self.render((element,))
        picture = deck.slides[0].shapes[0]
        self.assertGreater(picture.crop_left, 0)
        self.assertGreater(picture.crop_right, 0)
        self.assertAlmostEqual(picture.crop_top, 0, places=3)

    def test_chart_style_removes_default_grid_and_adds_labels(self):
        element = HybridElement(
            plan("chart", ElementRole.CHART, ProductionStrategy.NATIVE_CHART, "2024 12; 2025 18"),
            DEFAULT_LAYOUT_GRAMMAR.content_box(),
            chart=ChartSpec(("2024", "2025"), (ChartSeries("續約率", (12, 18)),)),
            chart_style=ChartStyle(show_gridlines=False, show_data_labels=True),
        )
        deck = self.render((element,))
        chart = deck.slides[0].shapes[0].chart
        self.assertFalse(chart.has_title)
        self.assertFalse(chart.value_axis.has_major_gridlines)
        self.assertFalse(chart.has_legend)
        self.assertTrue(chart.plots[0].has_data_labels)
        self.assertEqual(chart.plots[0].gap_width, 65)

    def test_shape_border_is_quiet_not_default_heavy(self):
        element = HybridElement(
            plan("card", ElementRole.SIMPLE_SHAPE, ProductionStrategy.NATIVE_SHAPE, "card"),
            DEFAULT_LAYOUT_GRAMMAR.columns(3)[0],
            shape_style=ShapeStyle("rounded_rectangle", "F4F7FB", "D6DFEA", 0.5),
        )
        deck = self.render((element,))
        self.assertAlmostEqual(deck.slides[0].shapes[0].line.width.pt, 0.5, places=1)

    def test_density_capacity_is_stricter_for_cjk(self):
        box = DEFAULT_LAYOUT_GRAMMAR.columns(3)[0]
        self.assertLess(max_characters_for_box(box, 17, cjk=True), max_characters_for_box(box, 17, cjk=False))

    def test_mixed_cjk_line_estimate_catches_orphan_risk_and_wider_repair(self):
        copy = "以可驗證的營運訊號，找出真正推動 Q3 表現的因素"
        box_type = type(DEFAULT_LAYOUT_GRAMMAR.content_box())
        self.assertEqual(estimated_line_count(copy, box_type(0.9, 3.1, 5.3, 0.7), 17), 2)
        self.assertEqual(estimated_line_count(copy, box_type(0.9, 3.1, 5.9, 0.7), 17), 1)

    def test_text_envelope_flags_line_and_font_departures(self):
        envelope = EditabilityEnvelope(
            expected_min_characters=1,
            expected_max_characters=48,
            expected_min_lines=1,
            expected_max_lines=1,
            minimum_font_size=18,
            maximum_font_size=24,
        )
        box = type(DEFAULT_LAYOUT_GRAMMAR.content_box())(0.9, 3.1, 5.3, 0.7)
        findings = audit_text_envelope(
            "subtitle", ElementRole.BODY, box,
            "以可驗證的營運訊號，找出真正推動 Q3 表現的因素", 17, envelope,
        )
        self.assertEqual(
            {finding.issue for finding in findings},
            {LayoutIssue.ENVELOPE_FONT_SIZE, LayoutIssue.ENVELOPE_LINE_COUNT},
        )

    def test_layout_audit_flags_long_title_and_weak_hierarchy(self):
        findings = audit_layout((
            ("title", ElementRole.TITLE, DEFAULT_LAYOUT_GRAMMAR.title_box(), "這是一個非常長而且沒有清楚焦點的報告式標題，超過簡報標題應有的長度與行數，並且繼續堆疊不必要的抽象說明", 18.0),
            ("body", ElementRole.BODY, DEFAULT_LAYOUT_GRAMMAR.content_box(), "Body", 17.0),
        ))
        issues = {finding.issue for finding in findings}
        self.assertIn(LayoutIssue.TITLE_LINE_COUNT, issues)
        self.assertIn(LayoutIssue.TYPOGRAPHY_HIERARCHY, issues)
        self.assertTrue(any(finding.severity is LayoutSeverity.REPAIR for finding in findings))

    def test_layout_audit_blocks_geometry_beyond_slide(self):
        box = type(DEFAULT_LAYOUT_GRAMMAR.content_box())(9.6, 1, 1, 1)
        findings = audit_layout((("body", ElementRole.BODY, box, "Text", 17.0),))
        self.assertEqual(findings[0].issue, LayoutIssue.OUT_OF_BOUNDS)
        self.assertEqual(findings[0].severity, LayoutSeverity.BLOCK)


if __name__ == "__main__":
    unittest.main()
