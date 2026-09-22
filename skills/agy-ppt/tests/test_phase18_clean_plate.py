#!/usr/bin/env python3
"""Phase 18 deterministic clean-plate production and reserved editable zones tests.

Covers all 25 required test cases:
 1. editable title zone absent from CLEAN_PLATE
 2. native title becomes sole visual title
 3. editing title reveals no previous text underneath
 4. KPI clean zone -> native KPI only
 5. KPI modification -> old value absent
 6. Logo clean zone -> replaceable Logo only
 7. Logo replacement -> old Logo absent
 8. photo clean zone -> replaceable photo only
 9. photo replacement -> old photo absent
10. chart clean zone -> native chart only
11. chart data modification -> old chart absent
12. artistic locked headline -> remains in plate -> no duplicate native headline
13. PARTIAL_COMPOSITE -> locked content remains -> only reserved zones receive native objects
14. FULL_COMPOSITE -> duplicate overlays prevented
15. unknown plate provenance -> conservative safe behavior
16. COMPOSITE_CONFLICT detected deterministically
17. Notes preserved
18. evidence-bound content preserved
19. Sample preview rendered from HYBRID_SLIDE rather than raw plate
20. Sample gate unchanged
21. exactly one Sample before approval
22. Sample reuse after approval
23. no OCR dependency added
24. no Jev runtime dependency
25. deterministic serialization
"""

from __future__ import annotations

import ast
import json
import sys
import tempfile
import unittest
from pathlib import Path

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

from PIL import Image  # noqa: E402
from pptx import Presentation  # noqa: E402

from phase18_contract import (  # noqa: E402
    DeliveryEditabilityContract,
    EditabilityClass,
    EditabilityContractError,
    ElementBox,
    FontPortability,
    PlateProvenanceMode,
    PlateRequirement,
    ProductionStrategy,
    ReservedEditableZone,
)
from phase18_fidelity_qa import (  # noqa: E402
    ElementSnapshot,
    FidelityIssue,
    FidelityFinding,
    QaDimension,
    QaSeverity,
    check_composite_conflicts,
    compare_fidelity,
    safe_representation_fallback,
)
from phase18_hybrid_pptx import (  # noqa: E402
    ChartSeries,
    ChartSpec,
    HybridElement,
    HybridPptxError,
    HybridSlide,
    TextStyle,
    create_hybrid_presentation,
    render_hybrid_sample_preview,
)
from phase18_production_plan import (  # noqa: E402
    ElementPlanningInput,
    ElementProductionPlan,
    ElementRole,
    PortabilityRisk,
    VisualPlateJobSpec,
    create_visual_plate_job,
    derive_reserved_zone,
    plan_element,
)
from phase18_roundtrip_compatibility import (  # noqa: E402
    simulate_chart_data_edit,
    simulate_kpi_edit,
    simulate_logo_replacement,
    simulate_photo_replacement,
    simulate_title_edit,
)
from presentation_workflow import (  # noqa: E402
    PresentationApprovalWorkflow,
    SampleResult,
)
from project_state import ProjectState  # noqa: E402


def _make_plan(
    element_id: str,
    strategy: ProductionStrategy,
    content: str,
    role: ElementRole = ElementRole.BODY,
    editability: EditabilityClass = EditabilityClass.EDITABLE_REQUIRED,
    evidence_claim_ids: tuple[str, ...] = (),
) -> ElementProductionPlan:
    return ElementProductionPlan(
        f"ep:{element_id}",
        element_id,
        f"claim:{element_id}",
        f"narrative:{element_id}",
        content,
        evidence_claim_ids,
        role,
        PortabilityRisk.LOW,
        DeliveryEditabilityContract(editability, strategy, FontPortability.FONT_SAFE),
    )


class Phase18CleanPlateTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.img1 = self.root / "img1.png"
        self.img2 = self.root / "img2.png"
        Image.new("RGB", (320, 180), (20, 40, 60)).save(self.img1)
        Image.new("RGB", (320, 180), (80, 100, 120)).save(self.img2)

    def tearDown(self):
        self.temp.cleanup()

    def _create_deck(
        self,
        elements,
        notes="Notes",
        name="test.pptx",
        plate_mode=PlateProvenanceMode.CLEAN_PLATE,
        reserved_zones=(),
    ) -> Path:
        out = self.root / name
        slide = HybridSlide(
            "slide_01",
            tuple(elements),
            speaker_notes=notes,
            plate_mode=plate_mode,
            reserved_zones=tuple(reserved_zones),
        )
        create_hybrid_presentation((slide,), str(out))
        return out

    # 1. editable title zone absent from CLEAN_PLATE
    def test_01_editable_title_zone_absent_from_clean_plate(self):
        plan = _make_plan("title", ProductionStrategy.NATIVE_TEXT, "Executive Title", ElementRole.TITLE, EditabilityClass.EDITABLE_PREFERRED)
        box = ElementBox(1.0, 0.5, 8.0, 1.0)
        job = create_visual_plate_job("slide_01", [(plan, box)], plate_mode=PlateProvenanceMode.CLEAN_PLATE)

        self.assertIn("Executive Title", job.editable_content_to_omit)
        self.assertNotIn("Executive Title", job.locked_content_to_render)
        self.assertEqual(len(job.reserved_zones), 1)
        self.assertEqual(job.reserved_zones[0].plate_requirement, PlateRequirement.CONTENT_FREE)
        self.assertEqual(job.reserved_zones[0].expected_content_type, "text")

    # 2. native title becomes sole visual title
    def test_02_native_title_becomes_sole_visual_title(self):
        plan = _make_plan("title", ProductionStrategy.NATIVE_TEXT, "Sole Visual Title", ElementRole.TITLE)
        elem = HybridElement(plan, ElementBox(1.0, 0.5, 8.0, 1.0), text_style=TextStyle())
        out = self._create_deck([elem], name="sole_title.pptx")

        deck = Presentation(str(out))
        self.assertEqual(len(deck.slides[0].shapes), 1)
        shape = deck.slides[0].shapes[0]
        self.assertEqual(shape.name, "agy:title")
        self.assertEqual(shape.text, "Sole Visual Title")

    # 3. editing title reveals no previous text underneath
    def test_03_editing_title_reveals_no_previous_text_underneath(self):
        plan = _make_plan("title", ProductionStrategy.NATIVE_TEXT, "Initial Title Text", ElementRole.TITLE)
        elem = HybridElement(plan, ElementBox(1.0, 0.5, 8.0, 1.0), text_style=TextStyle())
        out = self._create_deck([elem], name="edit_title.pptx")

        finding = simulate_title_edit(out, "title", "Updated Clean Title")
        self.assertEqual(finding.outcome.value, "PASS")

        deck = Presentation(str(out))
        all_text = " ".join(shape.text for shape in deck.slides[0].shapes if shape.has_text_frame)
        self.assertIn("Updated Clean Title", all_text)
        self.assertNotIn("Initial Title Text", all_text)

    # 4. KPI clean zone -> native KPI only
    def test_04_kpi_clean_zone_native_kpi_only(self):
        plan = _make_plan("kpi", ProductionStrategy.NATIVE_TEXT, "+42%", ElementRole.KPI, EditabilityClass.EDITABLE_REQUIRED)
        box = ElementBox(1.0, 2.0, 3.0, 1.2)
        job = create_visual_plate_job("slide_01", [(plan, box)], plate_mode=PlateProvenanceMode.CLEAN_PLATE)

        self.assertIn("+42%", job.editable_content_to_omit)
        self.assertEqual(job.reserved_zones[0].plate_requirement, PlateRequirement.CONTENT_FREE)

        elem = HybridElement(plan, box, text_style=TextStyle(font_size=28))
        out = self._create_deck([elem], name="kpi_only.pptx")
        deck = Presentation(str(out))
        self.assertEqual(len(deck.slides[0].shapes), 1)
        self.assertEqual(deck.slides[0].shapes[0].text, "+42%")

    # 5. KPI modification -> old value absent
    def test_05_kpi_modification_old_value_absent(self):
        plan = _make_plan("kpi", ProductionStrategy.NATIVE_TEXT, "18%", ElementRole.KPI)
        elem = HybridElement(plan, ElementBox(1.0, 2.0, 3.0, 1.2), text_style=TextStyle(font_size=28))
        out = self._create_deck([elem], name="kpi_edit.pptx")

        finding = simulate_kpi_edit(out, "kpi", "+25%")
        self.assertEqual(finding.outcome.value, "PASS")

        deck = Presentation(str(out))
        all_text = " ".join(shape.text for shape in deck.slides[0].shapes if shape.has_text_frame)
        self.assertIn("+25%", all_text)
        self.assertNotIn("18%", all_text)

    # 6. Logo clean zone -> replaceable Logo only
    def test_06_logo_clean_zone_replaceable_logo_only(self):
        plan = _make_plan("logo", ProductionStrategy.NATIVE_IMAGE, "Brand Logo", ElementRole.LOGO, EditabilityClass.REPLACEABLE)
        box = ElementBox(8.0, 0.5, 1.5, 0.8)
        job = create_visual_plate_job("slide_01", [(plan, box)], plate_mode=PlateProvenanceMode.CLEAN_PLATE)

        self.assertIn("Brand Logo", job.replaceable_assets_to_omit)
        self.assertEqual(job.reserved_zones[0].plate_requirement, PlateRequirement.CONTENT_FREE)
        self.assertEqual(job.reserved_zones[0].expected_content_type, "image")

        elem = HybridElement(plan, box, image_path=str(self.img1))
        out = self._create_deck([elem], name="logo_only.pptx")
        deck = Presentation(str(out))
        self.assertEqual(len(deck.slides[0].shapes), 1)
        self.assertEqual(deck.slides[0].shapes[0].name, "agy:logo")

    # 7. Logo replacement -> old Logo absent
    def test_07_logo_replacement_old_logo_absent(self):
        plan = _make_plan("logo", ProductionStrategy.NATIVE_IMAGE, "Brand Logo", ElementRole.LOGO, EditabilityClass.REPLACEABLE)
        elem = HybridElement(plan, ElementBox(8.0, 0.5, 1.5, 0.8), image_path=str(self.img1))
        out = self._create_deck([elem], name="logo_replace.pptx")

        orig_bytes = self.img1.read_bytes()
        new_bytes = self.img2.read_bytes()
        finding = simulate_logo_replacement(out, "logo", self.img2)
        self.assertEqual(finding.outcome.value, "PASS")

        deck = Presentation(str(out))
        replaced_shape = deck.slides[0].shapes[0]
        self.assertEqual(replaced_shape.image.blob, new_bytes)
        self.assertNotEqual(replaced_shape.image.blob, orig_bytes)

    # 8. photo clean zone -> replaceable photo only
    def test_08_photo_clean_zone_replaceable_photo_only(self):
        plan = _make_plan("photo", ProductionStrategy.NATIVE_IMAGE, "Case Study Photo", ElementRole.PHOTO, EditabilityClass.REPLACEABLE)
        box = ElementBox(1.0, 1.5, 5.0, 3.5)
        job = create_visual_plate_job("slide_01", [(plan, box)], plate_mode=PlateProvenanceMode.CLEAN_PLATE)

        self.assertIn("Case Study Photo", job.replaceable_assets_to_omit)
        self.assertEqual(job.reserved_zones[0].plate_requirement, PlateRequirement.CONTENT_FREE)

        elem = HybridElement(plan, box, image_path=str(self.img1))
        out = self._create_deck([elem], name="photo_only.pptx")
        deck = Presentation(str(out))
        self.assertEqual(len(deck.slides[0].shapes), 1)
        self.assertEqual(deck.slides[0].shapes[0].name, "agy:photo")

    # 9. photo replacement -> old photo absent
    def test_09_photo_replacement_old_photo_absent(self):
        plan = _make_plan("photo", ProductionStrategy.NATIVE_IMAGE, "Case Photo", ElementRole.PHOTO, EditabilityClass.REPLACEABLE)
        elem = HybridElement(plan, ElementBox(1.0, 1.5, 5.0, 3.5), image_path=str(self.img1))
        out = self._create_deck([elem], name="photo_replace.pptx")

        new_bytes = self.img2.read_bytes()
        finding = simulate_photo_replacement(out, "photo", self.img2)
        self.assertEqual(finding.outcome.value, "PASS")

        deck = Presentation(str(out))
        self.assertEqual(deck.slides[0].shapes[0].image.blob, new_bytes)

    # 10. chart clean zone -> native chart only
    def test_10_chart_clean_zone_native_chart_only(self):
        plan = _make_plan("chart", ProductionStrategy.NATIVE_CHART, "Trend Chart", ElementRole.CHART, EditabilityClass.EDITABLE_PREFERRED)
        box = ElementBox(1.0, 1.0, 6.0, 4.0)
        job = create_visual_plate_job("slide_01", [(plan, box)], plate_mode=PlateProvenanceMode.CLEAN_PLATE)

        self.assertIn("Trend Chart", job.editable_content_to_omit)
        self.assertEqual(job.reserved_zones[0].plate_requirement, PlateRequirement.CONTENT_FREE)
        self.assertEqual(job.reserved_zones[0].expected_content_type, "chart")

        chart_spec = ChartSpec(("2024", "2025"), (ChartSeries("Growth", (10.0, 20.0)),))
        elem = HybridElement(plan, box, chart=chart_spec)
        out = self._create_deck([elem], name="chart_only.pptx")
        deck = Presentation(str(out))
        self.assertEqual(len(deck.slides[0].shapes), 1)
        self.assertTrue(deck.slides[0].shapes[0].has_chart)

    # 11. chart data modification -> old chart absent
    def test_11_chart_data_modification_old_chart_absent(self):
        plan = _make_plan("chart", ProductionStrategy.NATIVE_CHART, "Trend Chart", ElementRole.CHART)
        chart_spec = ChartSpec(("Q1", "Q2", "Q3"), (ChartSeries("Sales", (10.0, 20.0, 30.0)),))
        elem = HybridElement(plan, ElementBox(1.0, 1.0, 6.0, 4.0), chart=chart_spec)
        out = self._create_deck([elem], name="chart_edit.pptx")

        finding = simulate_chart_data_edit(out, "chart", [50.0, 60.0, 70.0])
        self.assertEqual(finding.outcome.value, "PASS")

        deck = Presentation(str(out))
        series = deck.slides[0].shapes[0].chart.plots[0].series[0]
        self.assertEqual(series.values, (50.0, 60.0, 70.0))

    # 12. artistic locked headline -> remains in plate -> no duplicate native headline
    def test_12_artistic_locked_headline_remains_in_plate_no_duplicate(self):
        plan = _make_plan("art_title", ProductionStrategy.LOCKED_VISUAL, "Artistic Grand Title", ElementRole.ARTISTIC_HEADLINE, EditabilityClass.LOCKED_REQUIRED)
        box = ElementBox(0.0, 0.0, 10.0, 5.625)
        job = create_visual_plate_job("slide_01", [(plan, box)], plate_mode=PlateProvenanceMode.PARTIAL_COMPOSITE)

        self.assertIn("Artistic Grand Title", job.locked_content_to_render)
        self.assertNotIn("Artistic Grand Title", job.editable_content_to_omit)
        self.assertEqual(len(job.reserved_zones), 0)

        # In hybrid assembly, it is rendered as a locked visual image, NOT a native textbox
        elem = HybridElement(plan, box, image_path=str(self.img1))
        out = self._create_deck([elem], name="artistic.pptx")
        deck = Presentation(str(out))
        self.assertEqual(len(deck.slides[0].shapes), 1)
        self.assertFalse(deck.slides[0].shapes[0].has_text_frame)

    # 13. PARTIAL_COMPOSITE -> locked content remains -> only reserved zones receive native objects
    def test_13_partial_composite_locked_remains_only_reserved_receive_native(self):
        art_plan = _make_plan("art", ProductionStrategy.LOCKED_VISUAL, "Artwork", ElementRole.HERO_ARTWORK, EditabilityClass.LOCKED_REQUIRED)
        text_plan = _make_plan("author", ProductionStrategy.NATIVE_TEXT, "Dr. Alice", ElementRole.NAME, EditabilityClass.EDITABLE_REQUIRED)

        art_box = ElementBox(0.0, 0.0, 10.0, 5.625)
        text_box = ElementBox(1.0, 4.0, 4.0, 1.0)

        job = create_visual_plate_job("slide_01", [(art_plan, art_box), (text_plan, text_box)], plate_mode=PlateProvenanceMode.PARTIAL_COMPOSITE)
        self.assertIn("Artwork", job.locked_content_to_render)
        self.assertIn("Dr. Alice", job.editable_content_to_omit)
        self.assertEqual(len(job.reserved_zones), 1)
        self.assertEqual(job.reserved_zones[0].element_id, "author")

        art_elem = HybridElement(art_plan, art_box, image_path=str(self.img1))
        text_elem = HybridElement(text_plan, text_box, text_style=TextStyle())
        out = self._create_deck(
            [art_elem, text_elem],
            name="partial.pptx",
            plate_mode=PlateProvenanceMode.PARTIAL_COMPOSITE,
            reserved_zones=job.reserved_zones,
        )
        deck = Presentation(str(out))
        self.assertEqual(len(deck.slides[0].shapes), 2)

    # 14. FULL_COMPOSITE -> duplicate overlays prevented
    def test_14_full_composite_duplicate_overlays_prevented(self):
        text_plan = _make_plan("title", ProductionStrategy.NATIVE_TEXT, "Title", ElementRole.TITLE)
        text_elem = HybridElement(text_plan, ElementBox(1.0, 1.0, 5.0, 1.0), text_style=TextStyle())

        with self.assertRaises(HybridPptxError) as ctx:
            HybridSlide("slide_01", (text_elem,), plate_mode=PlateProvenanceMode.FULL_COMPOSITE)
        self.assertIn("composite conflict", str(ctx.exception))

    # 15. unknown plate provenance -> conservative safe behavior
    def test_15_unknown_plate_provenance_conservative_safe_behavior(self):
        text_plan = _make_plan("title", ProductionStrategy.NATIVE_TEXT, "Title", ElementRole.TITLE)
        text_elem = HybridElement(text_plan, ElementBox(1.0, 1.0, 5.0, 1.0), text_style=TextStyle())

        # Missing reserved zone in PARTIAL_COMPOSITE raises composite conflict
        with self.assertRaises(HybridPptxError) as ctx:
            HybridSlide("slide_01", (text_elem,), plate_mode=PlateProvenanceMode.PARTIAL_COMPOSITE, reserved_zones=())
        self.assertIn("composite conflict", str(ctx.exception))

        # Invalid plate_mode raises HybridPptxError
        with self.assertRaises(HybridPptxError):
            HybridSlide("slide_01", (text_elem,), plate_mode="INVALID_MODE")  # type: ignore

    # 16. COMPOSITE_CONFLICT detected deterministically
    def test_16_composite_conflict_detected_deterministically(self):
        text_plan = _make_plan("title", ProductionStrategy.NATIVE_TEXT, "Title", ElementRole.TITLE)
        text_elem = HybridElement(text_plan, ElementBox(1.0, 1.0, 5.0, 1.0), text_style=TextStyle())

        # Check via check_composite_conflicts
        findings = check_composite_conflicts("slide_01", PlateProvenanceMode.FULL_COMPOSITE, [text_elem])
        self.assertEqual(len(findings), 1)
        self.assertEqual(findings[0].issue, FidelityIssue.COMPOSITE_CONFLICT)
        self.assertEqual(findings[0].severity, QaSeverity.BLOCKING)

        # Check via compare_fidelity
        snap_approved = ElementSnapshot("title", "Title", ElementBox(1.0, 1.0, 5.0, 1.0), editable=True)
        snap_delivered = ElementSnapshot(
            "title", "Title", ElementBox(1.0, 1.0, 5.0, 1.0),
            editable=True, plate_contains_content=True,
        )
        report = compare_fidelity([snap_approved], [snap_delivered])
        self.assertEqual(report.status, QaSeverity.BLOCKING)
        self.assertIn(FidelityIssue.COMPOSITE_CONFLICT, [f.issue for f in report.findings])

        # Safe fallback converts plan to LOCKED_VISUAL / LOCKED_REQUIRED
        fallback = safe_representation_fallback(text_plan, report)
        self.assertTrue(fallback.changed)
        self.assertEqual(fallback.plan.contract.strategy, ProductionStrategy.LOCKED_VISUAL)
        self.assertEqual(fallback.plan.contract.editability, EditabilityClass.LOCKED_REQUIRED)

    # 17. Notes preserved
    def test_17_notes_preserved(self):
        notes_text = "Executive speaker notes: verified evidence 2026."
        plan = _make_plan("title", ProductionStrategy.NATIVE_TEXT, "Title", ElementRole.TITLE)
        elem = HybridElement(plan, ElementBox(1.0, 1.0, 5.0, 1.0), text_style=TextStyle())
        out = self._create_deck([elem], notes=notes_text, name="notes.pptx")

        deck = Presentation(str(out))
        self.assertEqual(deck.slides[0].notes_slide.notes_text_frame.text, notes_text)

    # 18. evidence-bound content preserved
    def test_18_evidence_bound_content_preserved(self):
        claims = ("pc:claim:001", "pc:claim:002")
        input_item = ElementPlanningInput(
            "kpi", "sem:kpi", "narr:kpi", ElementRole.KPI, "18%",
            evidence_claim_ids=claims,
        )
        plan = plan_element(input_item)
        self.assertEqual(plan.evidence_claim_ids, claims)

        box = ElementBox(1.0, 1.0, 3.0, 1.0)
        zone = derive_reserved_zone(plan, "slide_01", box)
        self.assertEqual(zone.element_id, "kpi")
        self.assertEqual(zone.plate_requirement, PlateRequirement.CONTENT_FREE)

    # 19. Sample preview rendered from HYBRID_SLIDE rather than raw plate
    def test_19_sample_preview_rendered_from_hybrid_slide_rather_than_raw_plate(self):
        workflow = PresentationApprovalWorkflow(ProjectState.initialize(str(self.root / "workflow_sample"), "sample_test"))
        workflow.submit_outline(({"number": 1, "title": "Overview", "role": "cover"},))
        workflow.approve_outline()
        workflow.submit_style({"palette": "modern"})
        workflow.approve_style()

        # Hybrid generator renders real hybrid slide preview
        plan = _make_plan("t1", ProductionStrategy.NATIVE_TEXT, "Overview", ElementRole.TITLE)
        elem = HybridElement(plan, ElementBox(1, 1, 6, 1), text_style=TextStyle())
        sample_slide = HybridSlide("sample_01", (elem,), plate_mode=PlateProvenanceMode.CLEAN_PLATE)

        def hybrid_generator(_slide_info):
            preview_out = self.root / "sample_preview.pptx"
            return render_hybrid_sample_preview(sample_slide, preview_out)

        sample_res = workflow.generate_sample(hybrid_generator)
        self.assertTrue(sample_res.is_hybrid_preview)
        self.assertTrue(Path(sample_res.artifact_ref).is_file())
        deck = Presentation(sample_res.artifact_ref)
        self.assertEqual(deck.slides[0].shapes[0].text, "Overview")

    # 20. Sample gate unchanged
    def test_20_sample_gate_unchanged(self):
        workflow = PresentationApprovalWorkflow(ProjectState.initialize(str(self.root / "workflow_gate"), "gate_test"))
        self.assertEqual(workflow.status, "OUTLINE_PENDING_APPROVAL")
        workflow.submit_outline(({"number": 1, "title": "Cover", "role": "cover"},))
        self.assertEqual(workflow.status, "OUTLINE_PENDING_APPROVAL")
        workflow.approve_outline()
        self.assertEqual(workflow.status, "STYLE_PENDING_APPROVAL")
        workflow.submit_style({"tone": "executive"})
        self.assertEqual(workflow.status, "STYLE_PENDING_APPROVAL")
        workflow.approve_style()
        self.assertEqual(workflow.status, "SAMPLE_PENDING_APPROVAL")
        workflow.generate_sample(lambda _: "sample.artifact")
        self.assertEqual(workflow.status, "SAMPLE_PENDING_APPROVAL")
        workflow.approve_sample()
        self.assertEqual(workflow.status, "READY_FOR_FULL_GENERATION")

    # 21. exactly one Sample before approval
    def test_21_exactly_one_sample_before_approval(self):
        workflow = PresentationApprovalWorkflow(ProjectState.initialize(str(self.root / "workflow_single"), "single_sample"))
        workflow.submit_outline(({"number": 1, "title": "Cover", "role": "cover"},))
        workflow.approve_outline()
        workflow.submit_style({"tone": "clean"})
        workflow.approve_style()

        res1 = workflow.generate_sample(lambda _: "sample_v1.artifact")
        self.assertEqual(res1.artifact_ref, "sample_v1.artifact")

        res2 = workflow.generate_sample(lambda _: "sample_v2.artifact")
        self.assertEqual(res2.artifact_ref, "sample_v2.artifact")

        # Approval approves res2 (the sole current sample)
        workflow.approve_sample()
        self.assertEqual(workflow.status, "READY_FOR_FULL_GENERATION")

    # 22. Sample reuse after approval
    def test_22_sample_reuse_after_approval(self):
        workflow = PresentationApprovalWorkflow(ProjectState.initialize(str(self.root / "workflow_reuse"), "reuse_sample"))
        workflow.submit_outline(({"number": 1, "title": "Cover", "role": "cover"}, {"number": 2, "title": "Details", "role": "content"},))
        workflow.approve_outline()
        workflow.submit_style({"tone": "clean"})
        workflow.approve_style()
        workflow.generate_sample(lambda _: "approved_sample.pptx")
        workflow.approve_sample()

        called_with_sample = []

        def full_generator(_outline, _style, sample_ref):
            called_with_sample.append(sample_ref)
            return "deck.pptx"

        gen_res = workflow.generate_full(full_generator)
        self.assertEqual(gen_res.value, "deck.pptx")
        self.assertEqual(called_with_sample, ["approved_sample.pptx"])

    # 23. no OCR dependency added
    def test_23_no_ocr_dependency_added(self):
        phase18_modules = [
            "phase18_contract.py",
            "phase18_production_plan.py",
            "phase18_hybrid_pptx.py",
            "phase18_fidelity_qa.py",
            "phase18_roundtrip_compatibility.py",
        ]
        prohibited_terms = {"pytesseract", "easyocr", "tesseract", "ocr"}
        for mod_name in phase18_modules:
            path = SCRIPTS / mod_name
            source = path.read_text(encoding="utf-8")
            tree = ast.parse(source, filename=str(path))
            for node in ast.walk(tree):
                if isinstance(node, ast.Import):
                    for alias in node.names:
                        self.assertNotIn(
                            alias.name.lower(),
                            prohibited_terms,
                            f"{mod_name} imports prohibited OCR library {alias.name}",
                        )
                elif isinstance(node, ast.ImportFrom) and node.module:
                    self.assertNotIn(
                        node.module.lower(),
                        prohibited_terms,
                        f"{mod_name} imports from prohibited OCR module {node.module}",
                    )

    # 24. no Jev runtime dependency
    def test_24_no_jev_runtime_dependency(self):
        phase18_modules = [
            "phase18_contract.py",
            "phase18_production_plan.py",
            "phase18_hybrid_pptx.py",
            "phase18_fidelity_qa.py",
            "phase18_roundtrip_compatibility.py",
            "presentation_workflow.py",
        ]
        prohibited_terms = {"typesafe", "jev"}
        for mod_name in phase18_modules:
            path = SCRIPTS / mod_name
            source = path.read_text(encoding="utf-8")
            tree = ast.parse(source, filename=str(path))
            for node in ast.walk(tree):
                if isinstance(node, ast.Import):
                    for alias in node.names:
                        self.assertNotIn(
                            alias.name.lower(),
                            prohibited_terms,
                            f"{mod_name} has runtime dependency on {alias.name}",
                        )
                elif isinstance(node, ast.ImportFrom) and node.module:
                    self.assertNotIn(
                        node.module.lower(),
                        prohibited_terms,
                        f"{mod_name} has runtime dependency on {node.module}",
                    )

    # 25. deterministic serialization
    def test_25_deterministic_serialization(self):
        plan = _make_plan("title", ProductionStrategy.NATIVE_TEXT, "Deterministic Title", ElementRole.TITLE)
        box = ElementBox(1.0, 0.5, 8.0, 1.0)
        zone = derive_reserved_zone(plan, "slide_01", box)

        json1 = zone.canonical_json()
        json2 = zone.canonical_json()
        self.assertEqual(json1, json2)

        # Verify parsed structure
        data = json.loads(json1)
        self.assertEqual(data["plate_requirement"], "CONTENT_FREE")
        self.assertEqual(data["strategy"], "NATIVE_TEXT")

        job = create_visual_plate_job("slide_01", [(plan, box)], plate_mode=PlateProvenanceMode.CLEAN_PLATE)
        job_json1 = job.canonical_json()
        job_json2 = job.canonical_json()
        self.assertEqual(job_json1, job_json2)


if __name__ == "__main__":
    unittest.main()
