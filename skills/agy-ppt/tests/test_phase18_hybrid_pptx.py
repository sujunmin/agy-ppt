#!/usr/bin/env python3
"""Phase 18.3 hybrid PowerPoint production tests."""

from __future__ import annotations

import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

from PIL import Image  # noqa: E402
from pptx import Presentation  # noqa: E402

from phase18_contract import DeliveryEditabilityContract, EditabilityClass, FontPortability, ProductionStrategy  # noqa: E402
from phase18_hybrid_pptx import (  # noqa: E402
    ChartSeries, ChartSpec, ElementBox, HybridElement, HybridPptxError,
    HybridSlide, ShapeStyle, TextStyle, create_hybrid_presentation,
)
from phase18_production_plan import ElementProductionPlan, ElementRole, PortabilityRisk  # noqa: E402
from presentation_workflow import PresentationApprovalWorkflow  # noqa: E402
from project_state import ProjectState  # noqa: E402


def manual_plan(element_id: str, strategy: ProductionStrategy, content: str, *, editability=EditabilityClass.EDITABLE_PREFERRED):
    return ElementProductionPlan(
        f"ep:{element_id}", element_id, f"semantic:{element_id}", "narrative:stable",
        content, ("pc:evidence",) if "18%" in content else (), ElementRole.KPI if "18%" in content else ElementRole.BODY,
        PortabilityRisk.LOW,
        DeliveryEditabilityContract(editability, strategy, FontPortability.FONT_SAFE),
    )


class Phase18HybridPptxTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.png = self.root / "visual.png"
        Image.new("RGB", (320, 180), (24, 64, 96)).save(self.png)

    def tearDown(self):
        self.temp.cleanup()

    def text_element(self, element_id="title", text="Editable title", **style):
        return HybridElement(
            manual_plan(element_id, ProductionStrategy.NATIVE_TEXT, text),
            ElementBox(0.7, 0.5, 8.6, 1.0),
            text_style=TextStyle(**style),
        )

    def render(self, elements, notes="Useful speaker notes", name="deck.pptx"):
        output = self.root / name
        self.assertTrue(create_hybrid_presentation((HybridSlide("slide_01", tuple(elements), notes),), str(output)))
        return output, Presentation(output)

    def test_native_business_text_is_editable_and_exact(self):
        _, deck = self.render((self.text_element(text="Name: Lin / KPI 18%"),))
        self.assertEqual(deck.slides[0].shapes[0].text, "Name: Lin / KPI 18%")

    def test_native_title_preserves_style(self):
        _, deck = self.render((self.text_element(font_name="Arial", font_size=30, bold=True, alignment="center"),))
        run = deck.slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
        self.assertTrue(run.font.bold)
        self.assertEqual(round(run.font.size.pt), 30)

    def test_native_body_preserves_paragraph_structure(self):
        _, deck = self.render((self.text_element(text="First\nSecond"),))
        paragraphs = deck.slides[0].shapes[0].text_frame.paragraphs
        self.assertEqual(tuple(item.text for item in paragraphs), ("First", "Second"))

    def test_native_simple_shape(self):
        element = HybridElement(
            manual_plan("card", ProductionStrategy.NATIVE_SHAPE, "card"), ElementBox(1, 1, 3, 2),
            shape_style=ShapeStyle("rounded_rectangle", "112233", "445566"),
        )
        _, deck = self.render((element,))
        self.assertEqual(deck.slides[0].shapes[0].name, "agy:card")

    def test_replaceable_logo_is_independent_picture(self):
        plan = manual_plan("logo", ProductionStrategy.NATIVE_IMAGE, "logo", editability=EditabilityClass.REPLACEABLE)
        _, deck = self.render((HybridElement(plan, ElementBox(8, 0.3, 1, 0.6), image_path=str(self.png)),))
        self.assertEqual(deck.slides[0].shapes[0].name, "agy:logo")

    def test_replaceable_photo_is_independent_picture(self):
        plan = manual_plan("photo", ProductionStrategy.NATIVE_IMAGE, "photo", editability=EditabilityClass.REPLACEABLE)
        _, deck = self.render((HybridElement(plan, ElementBox(5, 1, 4, 3), image_path=str(self.png)),))
        self.assertEqual(len(deck.slides[0].shapes), 1)

    def test_svg_uses_explicit_safe_fallback_when_library_cannot_embed_it(self):
        svg = self.root / "logo.svg"
        svg.write_text('<svg xmlns="http://www.w3.org/2000/svg" width="10" height="10"><rect width="10" height="10" fill="red"/></svg>', encoding="utf-8")
        plan = manual_plan("vector", ProductionStrategy.VECTOR_GRAPHIC, "vector", editability=EditabilityClass.REPLACEABLE)
        _, deck = self.render((HybridElement(plan, ElementBox(1, 1, 2, 2), image_path=str(svg), vector_fallback_path=str(self.png)),))
        self.assertIn("agy:vector", deck.slides[0].shapes[0].name)

    def test_native_structured_data_chart_preserves_values(self):
        plan = manual_plan("chart", ProductionStrategy.NATIVE_CHART, "2024 12; 2025 18")
        chart = ChartSpec(("2024", "2025"), (ChartSeries("Growth", (12, 18)),), "column")
        _, deck = self.render((HybridElement(plan, ElementBox(1, 1, 8, 3.5), chart=chart),))
        self.assertEqual(deck.slides[0].shapes[0].chart.series[0].values, (12.0, 18.0))

    def test_hybrid_editorial_chart_may_be_locked_region(self):
        plan = manual_plan("editorial-chart", ProductionStrategy.RASTER_REGION, "editorial chart", editability=EditabilityClass.LOCKED_PREFERRED)
        _, deck = self.render((HybridElement(plan, ElementBox(1, 1, 8, 3.5), image_path=str(self.png)),))
        self.assertEqual(deck.slides[0].shapes[0].name, "agy:editorial-chart")

    def test_raster_artwork_region_is_selective(self):
        plan = manual_plan("art", ProductionStrategy.LOCKED_VISUAL, "art", editability=EditabilityClass.LOCKED_REQUIRED)
        _, deck = self.render((self.text_element(), HybridElement(plan, ElementBox(5, 1.5, 4, 3), image_path=str(self.png))))
        self.assertEqual(len(deck.slides[0].shapes), 2)

    def test_mixed_native_and_raster_slide(self):
        plan = manual_plan("art", ProductionStrategy.RASTER_REGION, "art", editability=EditabilityClass.LOCKED_PREFERRED)
        _, deck = self.render((self.text_element(), HybridElement(plan, ElementBox(5, 1.5, 4, 3), image_path=str(self.png))))
        self.assertTrue(deck.slides[0].shapes[0].has_text_frame)
        self.assertFalse(deck.slides[0].shapes[1].has_text_frame)

    def test_full_raster_is_explicit_single_element_fallback(self):
        plan = manual_plan("full", ProductionStrategy.FULL_RASTER_SLIDE, "full", editability=EditabilityClass.LOCKED_REQUIRED)
        _, deck = self.render((HybridElement(plan, ElementBox(0, 0, 10, 5.625), image_path=str(self.png)),))
        self.assertEqual(len(deck.slides[0].shapes), 1)
        with self.assertRaises(HybridPptxError):
            HybridSlide("bad", (HybridElement(plan, ElementBox(0, 0, 10, 5.625), image_path=str(self.png)), self.text_element()))

    def test_traditional_chinese_text_is_exact(self):
        _, deck = self.render((self.text_element(text="內容有根據，重要內容可修改。"),))
        self.assertEqual(deck.slides[0].shapes[0].text, "內容有根據，重要內容可修改。")

    def test_english_text_is_exact(self):
        _, deck = self.render((self.text_element(text="Grounded and editable."),))
        self.assertEqual(deck.slides[0].shapes[0].text, "Grounded and editable.")

    def test_evidence_bound_kpi_remains_exact(self):
        element = HybridElement(manual_plan("kpi", ProductionStrategy.NATIVE_TEXT, "18%"), ElementBox(1, 1, 3, 1), text_style=TextStyle(font_size=34))
        _, deck = self.render((element,))
        self.assertEqual(deck.slides[0].shapes[0].text, "18%")

    def test_phase17_speaker_notes_are_preserved(self):
        notes = "Pause, explain the KPI, then transition to the recommendation."
        _, deck = self.render((self.text_element(),), notes=notes)
        self.assertIn(notes, deck.slides[0].notes_slide.notes_text_frame.text)

    def test_sample_approval_semantics_are_not_changed(self):
        workflow = PresentationApprovalWorkflow(ProjectState.initialize(str(self.root / "workflow"), "phase18"))
        outline = ({"number": 1, "title": "Opening", "role": "cover"},)
        workflow.submit_outline(outline)
        workflow.approve_outline()
        workflow.submit_style({"tone": "premium"})
        workflow.approve_style()
        workflow.generate_sample(lambda _: "sample-sha")
        self.assertEqual(workflow.approve_sample(), "READY_FOR_FULL_GENERATION")

    def test_pptx_opens_and_contains_expected_package_parts(self):
        output, deck = self.render((self.text_element(),))
        self.assertEqual(len(deck.slides), 1)
        with zipfile.ZipFile(output) as package:
            names = set(package.namelist())
        self.assertIn("ppt/slides/slide1.xml", names)
        self.assertIn("ppt/notesSlides/notesSlide1.xml", names)

    def test_object_order_and_identity_are_deterministic(self):
        shape_plan = manual_plan("card", ProductionStrategy.NATIVE_SHAPE, "card")
        elements = (self.text_element("title"), HybridElement(shape_plan, ElementBox(1, 2, 3, 2), shape_style=ShapeStyle()))
        _, first = self.render(elements, name="first.pptx")
        _, second = self.render(elements, name="second.pptx")
        self.assertEqual(tuple(shape.name for shape in first.slides[0].shapes), tuple(shape.name for shape in second.slides[0].shapes))


if __name__ == "__main__":
    unittest.main()
