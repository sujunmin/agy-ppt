#!/usr/bin/env python3
"""Phase 18.5 round-trip editability and compatibility QA tests."""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

from PIL import Image  # noqa: E402
from pptx import Presentation  # noqa: E402

from phase18_contract import (  # noqa: E402
    DeliveryEditabilityContract, EditabilityClass, FontPortability, ProductionStrategy,
)
from phase18_hybrid_pptx import (  # noqa: E402
    ChartSeries, ChartSpec, ElementBox, HybridElement, HybridSlide, TextStyle, create_hybrid_presentation,
)
from phase18_production_plan import ElementProductionPlan, ElementRole, PortabilityRisk  # noqa: E402
from phase18_roundtrip_compatibility import (  # noqa: E402
    QualificationType, RoundTripOutcome, audit_client_environments, run_roundtrip_qualification,
    simulate_chart_data_edit, simulate_kpi_edit, simulate_logo_replacement, simulate_photo_replacement,
    simulate_save_reopen, simulate_title_edit, verify_evidence_integrity_guard, verify_fallback_font,
    verify_table_boundary, verify_typography,
)


def make_plan(element_id: str, strategy: ProductionStrategy, content: str, role: ElementRole = ElementRole.BODY):
    return ElementProductionPlan(
        f"ep:{element_id}", element_id, f"claim:{element_id}", f"narrative:{element_id}",
        content, ("pc:evidence",) if "18%" in content else (), role,
        PortabilityRisk.LOW,
        DeliveryEditabilityContract(EditabilityClass.EDITABLE_REQUIRED, strategy, FontPortability.FONT_SAFE),
    )


class Phase18RoundTripCompatibilityTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.img1 = self.root / "img1.png"
        self.img2 = self.root / "img2.png"
        Image.new("RGB", (320, 180), (10, 20, 30)).save(self.img1)
        Image.new("RGB", (320, 180), (50, 60, 70)).save(self.img2)

    def tearDown(self):
        self.temp.cleanup()

    def _create_deck(self, elements, notes="Notes", name="test.pptx") -> Path:
        out = self.root / name
        slide = HybridSlide("slide_01", tuple(elements), speaker_notes=notes)
        create_hybrid_presentation((slide,), str(out))
        return out

    def test_kpi_edit_persists_within_envelope(self):
        elem = HybridElement(make_plan("kpi", ProductionStrategy.NATIVE_TEXT, "18%", ElementRole.KPI), ElementBox(1, 1, 3, 1), text_style=TextStyle())
        deck = self._create_deck([elem])
        finding = simulate_kpi_edit(deck, "kpi", "22%")
        self.assertEqual(finding.outcome, RoundTripOutcome.PASS)
        self.assertEqual(finding.qualification, QualificationType.STRUCTURAL_PROXY)
        reopened = Presentation(str(deck))
        self.assertEqual(reopened.slides[0].shapes[0].text, "22%")

    def test_title_edit_within_envelope_bounded_reflow(self):
        elem = HybridElement(make_plan("title", ProductionStrategy.NATIVE_TEXT, "Original Title", ElementRole.TITLE), ElementBox(1, 1, 6, 1), text_style=TextStyle())
        deck = self._create_deck([elem])
        finding = simulate_title_edit(deck, "title", "Updated Title Within Envelope", max_chars=40)
        self.assertEqual(finding.outcome, RoundTripOutcome.PASS)
        self.assertIn("stable reflow", finding.detail)

    def test_title_edit_beyond_envelope_emits_honest_warning(self):
        elem = HybridElement(make_plan("title", ProductionStrategy.NATIVE_TEXT, "Original Title", ElementRole.TITLE), ElementBox(1, 1, 6, 1), text_style=TextStyle())
        deck = self._create_deck([elem])
        long_title = "This is an extremely long title designed specifically to overflow the designated envelope bounds for slide headers"
        finding = simulate_title_edit(deck, "title", long_title, max_chars=40)
        self.assertEqual(finding.outcome, RoundTripOutcome.WARNING)
        self.assertIn("exceeds envelope", finding.detail)

    def test_logo_replacement_preserves_frame_and_placement(self):
        elem = HybridElement(make_plan("logo", ProductionStrategy.NATIVE_IMAGE, "logo"), ElementBox(8, 0.5, 2, 1), image_path=str(self.img1))
        deck = self._create_deck([elem])
        finding = simulate_logo_replacement(deck, "logo", self.img2)
        self.assertEqual(finding.outcome, RoundTripOutcome.PASS)
        reopened = Presentation(str(deck))
        self.assertEqual(reopened.slides[0].shapes[0].name, "agy:logo")
        self.assertEqual(reopened.slides[0].shapes[0].image.blob, self.img2.read_bytes())

    def test_photo_replacement_preserves_frame_and_reopens(self):
        elem = HybridElement(make_plan("photo", ProductionStrategy.NATIVE_IMAGE, "photo"), ElementBox(2, 2, 4, 3), image_path=str(self.img1))
        deck = self._create_deck([elem])
        finding = simulate_photo_replacement(deck, "photo", self.img2)
        self.assertEqual(finding.outcome, RoundTripOutcome.PASS)
        reopened = Presentation(str(deck))
        self.assertEqual(reopened.slides[0].shapes[0].image.blob, self.img2.read_bytes())

    def test_table_boundary_documented_honestly(self):
        finding = verify_table_boundary()
        self.assertEqual(finding.outcome, RoundTripOutcome.PASS)
        self.assertIn("not implemented", finding.detail)
        self.assertIn("LOCKED_VISUAL", finding.detail)

    def test_native_chart_data_modification_is_isolated_and_persists(self):
        chart_spec = ChartSpec(("Q1", "Q2", "Q3"), (ChartSeries("Sales", (10.0, 20.0, 30.0)),))
        elem = HybridElement(make_plan("chart", ProductionStrategy.NATIVE_CHART, "chart"), ElementBox(1, 1, 5, 4), chart=chart_spec)
        deck = self._create_deck([elem])
        finding = simulate_chart_data_edit(deck, "chart", (15.0, 25.0, 35.0), categories=("Q1", "Q2", "Q3"))
        self.assertEqual(finding.outcome, RoundTripOutcome.PASS)
        reopened = Presentation(str(deck))
        reopened_chart = reopened.slides[0].shapes[0].chart
        self.assertEqual(tuple(reopened_chart.series[0].values), (15.0, 25.0, 35.0))

    def test_save_reopen_preserves_package_and_structural_identity(self):
        elem1 = HybridElement(make_plan("title", ProductionStrategy.NATIVE_TEXT, "Title"), ElementBox(1, 1, 4, 1), text_style=TextStyle())
        elem2 = HybridElement(make_plan("logo", ProductionStrategy.NATIVE_IMAGE, "logo"), ElementBox(8, 1, 2, 1), image_path=str(self.img1))
        deck = self._create_deck([elem1, elem2], notes="Speaker notes text")
        out_deck = self.root / "saved.pptx"
        finding = simulate_save_reopen(deck, save_path=out_deck)
        self.assertEqual(finding.outcome, RoundTripOutcome.PASS)
        reopened = Presentation(str(out_deck))
        self.assertEqual(len(reopened.slides), 1)
        self.assertEqual([s.name for s in reopened.slides[0].shapes], ["agy:title", "agy:logo"])
        self.assertEqual(reopened.slides[0].notes_slide.notes_text_frame.text, "Speaker notes text")

    def test_traditional_chinese_typography(self):
        text = "繁體中文測試：年增率 18.5%，毛利率穩健成長。"
        finding = verify_typography(text, language="zh-TW", font_name="Microsoft JhengHei")
        self.assertEqual(finding.outcome, RoundTripOutcome.PASS)
        self.assertIn("ZH-TW", finding.scenario)

    def test_english_typography(self):
        text = "Q3 Performance Review: Revenue up 18.5% YoY with steady operational margin."
        finding = verify_typography(text, language="en", font_name="Arial")
        self.assertEqual(finding.outcome, RoundTripOutcome.PASS)
        self.assertIn("EN", finding.scenario)

    def test_fallback_font_behavior(self):
        safe_finding = verify_fallback_font("Arial", FontPortability.FONT_SAFE)
        self.assertEqual(safe_finding.outcome, RoundTripOutcome.PASS)
        warning_finding = verify_fallback_font("CustomExoticFont", FontPortability.FONT_CRITICAL)
        self.assertEqual(warning_finding.outcome, RoundTripOutcome.WARNING)
        self.assertIn("protected visual", warning_finding.detail)

    def test_evidence_integrity_distinguishes_intentional_from_accidental(self):
        intentional = verify_evidence_integrity_guard("18%", "22%", is_intentional_user_edit=True)
        self.assertEqual(intentional.outcome, RoundTripOutcome.PASS)
        self.assertIn("intentional user edit", intentional.detail)

        accidental = verify_evidence_integrity_guard("18%", "1.8%", is_intentional_user_edit=False)
        self.assertEqual(accidental.outcome, RoundTripOutcome.BLOCKING)
        self.assertIn("unauthorized pipeline mutation", accidental.detail)

    def test_actual_client_vs_structural_proxy_reporting(self):
        audit = audit_client_environments()
        self.assertIn("unexecuted_environments", audit)
        self.assertIn("Windows Microsoft PowerPoint", audit["unexecuted_environments"])
        self.assertEqual(audit["qualification_type"], QualificationType.STRUCTURAL_PROXY)

        report = run_roundtrip_qualification(self.root)
        self.assertEqual(report.qualification, QualificationType.STRUCTURAL_PROXY)
        self.assertTrue(len(report.executed_scenarios) >= 12)
        self.assertIn("Windows Microsoft PowerPoint", report.unexecuted_environments)
        json_report = report.canonical_json()
        self.assertIn("STRUCTURAL_PROXY", json_report)
        self.assertIn("unexecuted_environments", json_report)

    def test_roundtrip_report_counts_and_canonical_json(self):
        report = run_roundtrip_qualification(self.root)
        self.assertEqual(report.count(RoundTripOutcome.PASS) + report.count(RoundTripOutcome.WARNING) + report.count(RoundTripOutcome.REVIEW_REQUIRED) + report.count(RoundTripOutcome.BLOCKING), len(report.findings))
        canonical = report.canonical_json()
        self.assertEqual(canonical, report.canonical_json())

    def test_kpi_missing_shape_reports_blocking(self):
        elem = HybridElement(make_plan("other", ProductionStrategy.NATIVE_TEXT, "Other"), ElementBox(1, 1, 2, 1), text_style=TextStyle())
        deck = self._create_deck([elem])
        finding = simulate_kpi_edit(deck, "missing_kpi", "30%")
        self.assertEqual(finding.outcome, RoundTripOutcome.BLOCKING)
        self.assertIn("not found", finding.detail)


if __name__ == "__main__":
    unittest.main()
