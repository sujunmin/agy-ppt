#!/usr/bin/env python3
"""Phase 18.3 conservative hybrid PowerPoint production.

The renderer consumes an approved element production plan.  It never chooses
evidence, wording, or narrative meaning; unsupported native representations
fall back to the explicitly supplied locked visual.
"""

from __future__ import annotations

import hashlib
import json
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Iterable

from phase18_contract import (
    ElementBox,
    PlateProvenanceMode,
    PlateRequirement,
    ProductionStrategy,
    ReservedEditableZone,
)
from phase18_production_plan import ElementProductionPlan
from presentation_workflow import SampleArtifact


ERROR_HYBRID_PPTX_INVALID = "PHASE18_HYBRID_PPTX_INVALID"


class HybridPptxError(Exception):
    def __init__(self, error_code: str, message: str) -> None:
        super().__init__(message)
        self.error_code = error_code


@dataclass(frozen=True)
class TextStyle:
    font_name: str = "Arial"
    font_size: float = 20.0
    bold: bool = False
    color: str = "000000"
    alignment: str = "left"

    def __post_init__(self) -> None:
        if not isinstance(self.font_name, str) or not self.font_name.strip() or self.font_size <= 0:
            raise HybridPptxError(ERROR_HYBRID_PPTX_INVALID, "text style is invalid")
        if len(self.color) != 6 or any(char not in "0123456789ABCDEFabcdef" for char in self.color):
            raise HybridPptxError(ERROR_HYBRID_PPTX_INVALID, "text color must be six hex digits")
        if self.alignment not in {"left", "center", "right"}:
            raise HybridPptxError(ERROR_HYBRID_PPTX_INVALID, "text alignment is invalid")


@dataclass(frozen=True)
class ShapeStyle:
    kind: str = "rectangle"
    fill_color: str = "FFFFFF"
    line_color: str = "FFFFFF"

    def __post_init__(self) -> None:
        if self.kind not in {"rectangle", "rounded_rectangle", "divider", "right_arrow"}:
            raise HybridPptxError(ERROR_HYBRID_PPTX_INVALID, "native shape kind is unsupported")
        for color in (self.fill_color, self.line_color):
            if len(color) != 6 or any(char not in "0123456789ABCDEFabcdef" for char in color):
                raise HybridPptxError(ERROR_HYBRID_PPTX_INVALID, "shape color is invalid")


@dataclass(frozen=True)
class ChartSeries:
    name: str
    values: tuple[float, ...]

    def __post_init__(self) -> None:
        if not isinstance(self.name, str) or not self.name.strip() or not self.values:
            raise HybridPptxError(ERROR_HYBRID_PPTX_INVALID, "chart series is invalid")
        if any(not isinstance(value, (int, float)) or isinstance(value, bool) for value in self.values):
            raise HybridPptxError(ERROR_HYBRID_PPTX_INVALID, "chart values must be numeric")


@dataclass(frozen=True)
class ChartSpec:
    categories: tuple[str, ...]
    series: tuple[ChartSeries, ...]
    kind: str = "column"

    def __post_init__(self) -> None:
        if not self.categories or any(not isinstance(item, str) or not item for item in self.categories):
            raise HybridPptxError(ERROR_HYBRID_PPTX_INVALID, "chart categories are required")
        if not self.series or any(not isinstance(item, ChartSeries) for item in self.series):
            raise HybridPptxError(ERROR_HYBRID_PPTX_INVALID, "chart series are required")
        if any(len(item.values) != len(self.categories) for item in self.series):
            raise HybridPptxError(ERROR_HYBRID_PPTX_INVALID, "chart values must match categories")
        if self.kind not in {"column", "bar", "line"}:
            raise HybridPptxError(ERROR_HYBRID_PPTX_INVALID, "chart kind is unsupported")


@dataclass(frozen=True)
class HybridElement:
    plan: ElementProductionPlan
    box: ElementBox
    text_style: TextStyle | None = None
    shape_style: ShapeStyle | None = None
    image_path: str | None = None
    vector_fallback_path: str | None = None
    chart: ChartSpec | None = None

    def __post_init__(self) -> None:
        if not isinstance(self.plan, ElementProductionPlan) or not isinstance(self.box, ElementBox):
            raise HybridPptxError(ERROR_HYBRID_PPTX_INVALID, "approved element plan and box are required")
        strategy = self.plan.contract.strategy
        if strategy is ProductionStrategy.NATIVE_TEXT and not isinstance(self.text_style, TextStyle):
            raise HybridPptxError(ERROR_HYBRID_PPTX_INVALID, "native text requires text style")
        if strategy is ProductionStrategy.NATIVE_SHAPE and not isinstance(self.shape_style, ShapeStyle):
            raise HybridPptxError(ERROR_HYBRID_PPTX_INVALID, "native shape requires shape style")
        if strategy in {ProductionStrategy.NATIVE_IMAGE, ProductionStrategy.RASTER_REGION, ProductionStrategy.LOCKED_VISUAL, ProductionStrategy.FULL_RASTER_SLIDE}:
            if not self.image_path or not Path(self.image_path).is_file():
                raise HybridPptxError(ERROR_HYBRID_PPTX_INVALID, "image-backed strategy requires an existing image")
        if strategy is ProductionStrategy.VECTOR_GRAPHIC and not self.image_path:
            raise HybridPptxError(ERROR_HYBRID_PPTX_INVALID, "vector strategy requires a vector source")
        if strategy is ProductionStrategy.NATIVE_CHART and not isinstance(self.chart, ChartSpec):
            raise HybridPptxError(ERROR_HYBRID_PPTX_INVALID, "native chart requires structured data")


@dataclass(frozen=True)
class HybridSlide:
    slide_id: str
    elements: tuple[HybridElement, ...]
    speaker_notes: str = ""
    background_color: str = "FFFFFF"
    plate_mode: PlateProvenanceMode = PlateProvenanceMode.CLEAN_PLATE
    reserved_zones: tuple[ReservedEditableZone, ...] = ()

    def __post_init__(self) -> None:
        if not isinstance(self.slide_id, str) or not self.slide_id.strip():
            raise HybridPptxError(ERROR_HYBRID_PPTX_INVALID, "slide id is required")
        elements = tuple(self.elements)
        if not elements or any(not isinstance(item, HybridElement) for item in elements):
            raise HybridPptxError(ERROR_HYBRID_PPTX_INVALID, "slide elements are required")
        if len({item.plan.element_id for item in elements}) != len(elements):
            raise HybridPptxError(ERROR_HYBRID_PPTX_INVALID, "element ids must be unique per slide")
        full = [item for item in elements if item.plan.contract.strategy is ProductionStrategy.FULL_RASTER_SLIDE]
        if full and len(elements) != 1:
            raise HybridPptxError(ERROR_HYBRID_PPTX_INVALID, "full-raster fallback cannot hide other planned objects")
        if len(self.background_color) != 6:
            raise HybridPptxError(ERROR_HYBRID_PPTX_INVALID, "background color is invalid")
        if not isinstance(self.plate_mode, PlateProvenanceMode):
            raise HybridPptxError(ERROR_HYBRID_PPTX_INVALID, "plate mode is invalid")
        zones = tuple(self.reserved_zones)
        if any(not isinstance(z, ReservedEditableZone) for z in zones):
            raise HybridPptxError(ERROR_HYBRID_PPTX_INVALID, "reserved zones are invalid")

        # Composite conflict enforcement:
        # 1. FULL_COMPOSITE prohibits native editable overlays
        if self.plate_mode is PlateProvenanceMode.FULL_COMPOSITE:
            for item in elements:
                if item.plan.contract.strategy in {
                    ProductionStrategy.NATIVE_TEXT,
                    ProductionStrategy.NATIVE_IMAGE,
                    ProductionStrategy.NATIVE_CHART,
                }:
                    raise HybridPptxError(
                        ERROR_HYBRID_PPTX_INVALID,
                        f"composite conflict: cannot overlay native element '{item.plan.element_id}' on FULL_COMPOSITE plate",
                    )
        # 2. PARTIAL_COMPOSITE requires a CONTENT_FREE reserved zone for each native overlay
        elif self.plate_mode is PlateProvenanceMode.PARTIAL_COMPOSITE:
            content_free_ids = {
                z.element_id for z in zones if z.plate_requirement is PlateRequirement.CONTENT_FREE
            }
            for item in elements:
                if item.plan.contract.strategy in {
                    ProductionStrategy.NATIVE_TEXT,
                    ProductionStrategy.NATIVE_IMAGE,
                    ProductionStrategy.NATIVE_CHART,
                }:
                    if item.plan.element_id not in content_free_ids:
                        raise HybridPptxError(
                            ERROR_HYBRID_PPTX_INVALID,
                            f"composite conflict: native element '{item.plan.element_id}' lacks CONTENT_FREE reserved zone on PARTIAL_COMPOSITE plate",
                        )

        object.__setattr__(self, "elements", elements)
        object.__setattr__(self, "reserved_zones", zones)


def _shape_type(kind: str):
    from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
    return {
        "rectangle": MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        "rounded_rectangle": MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        "divider": MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        "right_arrow": MSO_AUTO_SHAPE_TYPE.RIGHT_ARROW,
    }[kind]


def _set_object_name(shape, element: HybridElement, suffix: str = "") -> None:
    shape.name = f"agy:{element.plan.element_id}{suffix}"


def _add_element(slide, element: HybridElement) -> None:
    from pptx.chart.data import ChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    box = element.box
    geometry = (Inches(box.left), Inches(box.top), Inches(box.width), Inches(box.height))
    strategy = element.plan.contract.strategy
    if strategy is ProductionStrategy.NATIVE_TEXT:
        shape = slide.shapes.add_textbox(*geometry)
        frame = shape.text_frame
        frame.clear()
        lines = element.plan.approved_content.splitlines() or [element.plan.approved_content]
        frame.paragraphs[0].text = lines[0]
        for line in lines[1:]:
            frame.add_paragraph().text = line
        alignments = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
        for paragraph in frame.paragraphs:
            paragraph.alignment = alignments[element.text_style.alignment]
            for run in paragraph.runs:
                run.font.name = element.text_style.font_name
                run.font.size = Pt(element.text_style.font_size)
                run.font.bold = element.text_style.bold
                run.font.color.rgb = RGBColor.from_string(element.text_style.color)
        _set_object_name(shape, element)
    elif strategy is ProductionStrategy.NATIVE_SHAPE:
        shape = slide.shapes.add_shape(_shape_type(element.shape_style.kind), *geometry)
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(element.shape_style.fill_color)
        shape.line.color.rgb = RGBColor.from_string(element.shape_style.line_color)
        _set_object_name(shape, element)
    elif strategy is ProductionStrategy.NATIVE_CHART:
        data = ChartData()
        data.categories = element.chart.categories
        for series in element.chart.series:
            data.add_series(series.name, series.values)
        chart_types = {
            "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
            "bar": XL_CHART_TYPE.BAR_CLUSTERED,
            "line": XL_CHART_TYPE.LINE,
        }
        graphic_frame = slide.shapes.add_chart(chart_types[element.chart.kind], *geometry, data)
        graphic_frame.chart.has_legend = len(element.chart.series) > 1
        _set_object_name(graphic_frame, element)
    else:
        source = element.image_path
        suffix = ""
        if strategy is ProductionStrategy.VECTOR_GRAPHIC:
            try:
                shape = slide.shapes.add_picture(source, *geometry)
                _set_object_name(shape, element)
                return
            except Exception as exc:
                fallback = element.vector_fallback_path
                if not fallback or not Path(fallback).is_file():
                    raise HybridPptxError(ERROR_HYBRID_PPTX_INVALID, "vector source needs a compatible raster fallback") from exc
                source = fallback
                suffix = ":vector-fallback"
        shape = slide.shapes.add_picture(source, *geometry)
        _set_object_name(shape, element, suffix)


def create_hybrid_presentation(
    slides: Iterable[HybridSlide],
    output_path: str,
    *,
    aspect_ratio: str = "16:9",
) -> bool:
    """Create a hybrid PPTX without changing approved semantic content."""
    clean = tuple(slides)
    if not clean or any(not isinstance(item, HybridSlide) for item in clean):
        raise HybridPptxError(ERROR_HYBRID_PPTX_INVALID, "hybrid slides are required")
    if len({item.slide_id for item in clean}) != len(clean):
        raise HybridPptxError(ERROR_HYBRID_PPTX_INVALID, "slide ids must be unique")
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.util import Inches

        prs = Presentation()
        if aspect_ratio == "16:9":
            prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)
        elif aspect_ratio == "4:3":
            prs.slide_width, prs.slide_height = Inches(10), Inches(7.5)
        else:
            raise HybridPptxError(ERROR_HYBRID_PPTX_INVALID, "unsupported aspect ratio")
        for spec in clean:
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            background = slide.background.fill
            background.solid()
            background.fore_color.rgb = RGBColor.from_string(spec.background_color)
            for element in spec.elements:
                _add_element(slide, element)
            if spec.speaker_notes:
                notes = slide.notes_slide.notes_text_frame
                notes.clear()
                notes.text = spec.speaker_notes
        destination = Path(output_path)
        destination.parent.mkdir(parents=True, exist_ok=True)
        prs.save(destination)
        return True
    except HybridPptxError:
        raise
    except Exception as exc:
        raise HybridPptxError(ERROR_HYBRID_PPTX_INVALID, "hybrid PowerPoint production failed") from exc


def render_hybrid_sample_preview(
    slide: HybridSlide,
    output_path: str | Path,
    *,
    aspect_ratio: str = "16:9",
) -> SampleArtifact:
    """Render a deterministic hybrid sample preview artifact from a HybridSlide specification."""
    dest = Path(output_path)
    dest.parent.mkdir(parents=True, exist_ok=True)
    create_hybrid_presentation((slide,), str(dest), aspect_ratio=aspect_ratio)
    manifest = {
        "slide_id": slide.slide_id,
        "speaker_notes": slide.speaker_notes,
        "background_color": slide.background_color,
        "plate_mode": slide.plate_mode.value,
        "reserved_zones": [zone.internal_dict() for zone in slide.reserved_zones],
        "elements": [
            {
                "plan": element.plan.internal_dict(),
                "box": element.box.internal_dict(),
                "text_style": asdict(element.text_style) if element.text_style else None,
                "shape_style": asdict(element.shape_style) if element.shape_style else None,
                "chart": asdict(element.chart) if element.chart else None,
                "image_sha256": (
                    hashlib.sha256(Path(element.image_path).read_bytes()).hexdigest()
                    if element.image_path and Path(element.image_path).is_file()
                    else None
                ),
            }
            for element in slide.elements
        ],
    }
    manifest_hash = hashlib.sha256(
        json.dumps(manifest, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    ).hexdigest()
    preview_hash = hashlib.sha256(dest.read_bytes()).hexdigest()
    return SampleArtifact(
        artifact_ref=str(dest),
        provenance={
            "artifact_kind": "HYBRID_PREVIEW",
            "artifact_ref": str(dest),
            "hybrid_manifest_sha256": manifest_hash,
            "rendered_preview_sha256": preview_hash,
        },
    )


__all__ = [
    "ERROR_HYBRID_PPTX_INVALID", "ChartSeries", "ChartSpec", "ElementBox",
    "HybridElement", "HybridPptxError", "HybridSlide", "ShapeStyle", "TextStyle",
    "create_hybrid_presentation", "render_hybrid_sample_preview",
]
